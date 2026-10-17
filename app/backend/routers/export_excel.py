import io

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lib.dates import today_iso
from lib.auth import require_auth
from lib.db import db
from routers.alokasi import compute_hitung
from routers.rekap import build_rekap

router = APIRouter(dependencies=[Depends(require_auth)])

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill(fill_type="solid", start_color="065F46", end_color="065F46")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


@router.get("/export/excel")
async def export_excel():
    rekap = await build_rekap()
    hitung = await compute_hitung()
    desas = await db.desa.find({}, {"_id": 0}).to_list(500)
    persons = await db.persons.find({}, {"_id": 0}).sort("created_at", 1).to_list(100000)
    desa_map = {d["id"]: d for d in desas}
    kecs = {k.id: k.nama for k in rekap.kecamatan}

    wb = Workbook()
    wb.remove(wb.active)

    rows = []
    for i, k in enumerate(rekap.kecamatan, start=1):
        rows.append([i, k.nama, k.desa_count, k.desa_terisi, k.perangkat, k.penggiat, k.total_orang, k.jiwa, k.kk])
    rows.append(["", "TOTAL KABUPATEN", rekap.desa_count, "", rekap.total.perangkat, rekap.total.penggiat,
                 rekap.total.total_orang, rekap.total.jiwa, rekap.total.kk])
    _sheet(wb, "Rekap Kecamatan",
           ["No", "Kecamatan", "Jml Desa", "Desa Terisi", "Perangkat", "Penggiat", "Total Orang", "Anggota KK (Jiwa)", "KK"],
           rows, [5, 24, 10, 12, 11, 11, 12, 18, 8])

    rows = []
    no = 1
    for k in rekap.kecamatan:
        for d in k.desa:
            rows.append([no, k.nama, d.nama, d.perangkat, d.penggiat, d.total_orang, d.jiwa, d.kk])
            no += 1
    _sheet(wb, "Rekap Desa",
           ["No", "Kecamatan", "Desa", "Perangkat", "Penggiat", "Total Orang", "Anggota KK (Jiwa)", "KK"],
           rows, [5, 24, 26, 11, 11, 12, 18, 8])

    rows = []
    for i, p in enumerate(persons, start=1):
        d = desa_map.get(p.get("desa_id", ""), {})
        rows.append([i, p.get("nama", ""), p.get("nik", ""), p.get("no_kk", ""),
                     "Perangkat" if p.get("jenis") == "perangkat" else "Penggiat",
                     p.get("jabatan", ""), kecs.get(p.get("kecamatan_id", ""), ""), d.get("nama", ""),
                     p.get("jumlah_anggota_kk", 0)])
    _sheet(wb, "Detail Perangkat & Penggiat",
           ["No", "Nama", "NIK", "No. KK", "Jenis", "Jabatan", "Kecamatan", "Desa", "Jml Anggota KK"],
           rows, [5, 30, 20, 20, 11, 26, 20, 24, 16])

    rows = []
    no = 1
    for p in persons:
        d = desa_map.get(p.get("desa_id", ""), {})
        for a in p.get("anggota", []):
            rows.append([no, p.get("nama", ""), d.get("nama", ""), a.get("nik", ""), a.get("nama", ""),
                         a.get("hubungan", ""), a.get("jenis_kelamin", ""), a.get("tgl_lahir", ""),
                         a.get("pendidikan", "")])
            no += 1
    _sheet(wb, "Anggota Keluarga",
           ["No", "Kepala Data (Perangkat/Penggiat)", "Desa", "NIK Anggota", "Nama Anggota",
            "Hubungan", "Jenis Kelamin", "Tgl Lahir", "Pendidikan"],
           rows, [5, 30, 22, 20, 30, 16, 14, 12, 18])

    rows = []
    for i, it in enumerate(hitung.items, start=1):
        rows.append([i, it.nama, it.target, it.qty, it.rate, it.frekuensi, it.jumlah, it.keterangan])
    rows.append(["", "TOTAL ALOKASI", "", "", "", "", hitung.total, ""])
    _sheet(wb, "Alokasi Dana",
           ["No", "Nama Item", "Dasar Hitung", "Jumlah (Qty)", "Tarif (Rp)", "Frekuensi (x/tahun)",
            "Jumlah (Rp)", "Keterangan"],
           rows, [5, 40, 14, 14, 16, 18, 20, 30])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    today = today_iso("Asia/Jakarta")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="rekap_perangkat_penggiat_batanghari_{today}.xlsx"'},
    )
